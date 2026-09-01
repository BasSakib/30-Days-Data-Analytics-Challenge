import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

df = pd.read_csv("/home/claude/30-day-data-journey/data/sales_data_raw.csv")

wb = Workbook()
ws = wb.active
ws.title = "Sales"

# Write header + data
headers = list(df.columns)
ws.append(headers)
for row in df.itertuples(index=False):
    ws.append(list(row))

# Header formatting
header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(name="Arial", bold=True, color="FFFFFF")
for c in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=c)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center")

# Body font
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
    for cell in row:
        cell.font = Font(name="Arial", size=10)

# Number formats: OrderDate stays text here (raw/mixed formats — this is
# the exercise), UnitPrice/Revenue as currency, Quantity as number
unitprice_col = headers.index("UnitPrice") + 1
revenue_col = headers.index("Revenue") + 1
qty_col = headers.index("Quantity") + 1

for r in range(2, ws.max_row + 1):
    up = ws.cell(row=r, column=unitprice_col)
    if up.value not in (None, ""):
        up.number_format = '"$"#,##0.00'
    rv = ws.cell(row=r, column=revenue_col)
    if rv.value not in (None, ""):
        rv.number_format = '"$"#,##0.00'
    q = ws.cell(row=r, column=qty_col)
    if q.value not in (None, ""):
        q.number_format = '0'
        q.alignment = Alignment(horizontal="right")

# Auto-fit column widths (approx, based on content length)
for c in range(1, len(headers) + 1):
    col_letter = get_column_letter(c)
    max_len = max(
        [len(str(ws.cell(row=r, column=c).value)) for r in range(1, ws.max_row + 1)]
    )
    ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 30)

# Freeze header row
ws.freeze_panes = "A2"

# Make it a proper Excel Table
last_col_letter = get_column_letter(len(headers))
table_ref = f"A1:{last_col_letter}{ws.max_row}"
tbl = Table(displayName="tbl_Sales", ref=table_ref)
tbl.tableStyleInfo = TableStyleInfo(
    name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False
)
ws.add_table(tbl)

# Notes sheet
notes = wb.create_sheet("Notes")
notes["A1"] = "Data Quality Observations — Day 01"
notes["A1"].font = Font(name="Arial", bold=True, size=12)
observations = [
    "",
    "1. Missing values found in: Region, SalesRep, UnitPrice, Quantity columns.",
    "2. OrderDate is inconsistently formatted — mix of YYYY-MM-DD, DD/MM/YYYY, and MM-DD-YYYY.",
    "3. Region and Channel have inconsistent casing (e.g. 'NORTH', 'north', 'North') and extra whitespace.",
    "4. Duplicate OrderID values are present (~40 rows) — same order appears more than once.",
    "5. These issues are intentional and will be cleaned in Day 04 (Excel) and Day 09 (Power Query).",
]
for i, line in enumerate(observations, start=2):
    notes.cell(row=i, column=1, value=line).font = Font(name="Arial", size=11)
notes.column_dimensions["A"].width = 100

wb.save("/home/claude/30-day-data-journey/phase1-excel/day01/day01_solved.xlsx")
print("Day 01 saved.")
