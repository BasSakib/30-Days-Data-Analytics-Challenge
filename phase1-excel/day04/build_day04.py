import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

raw = pd.read_csv("/home/claude/30-day-data-journey/data/sales_data_raw.csv")
raw_dupes = int(raw.duplicated(subset=["OrderID"]).sum())
raw_missing = raw.isnull().sum()

# --- Perform the cleaning in Python (mirrors the manual Excel steps) ---
df = raw.drop_duplicates(subset=["OrderID"]).copy()
rows_removed = len(raw) - len(df)

df["Region"] = df["Region"].fillna("Unknown").astype(str).str.strip()
df["Region"] = df["Region"].where(df["Region"] == "", df["Region"]).replace("", "Unknown")
df["Region"] = df["Region"].apply(lambda x: x.title() if x != "Unknown" else x)

df["SalesRep"] = df["SalesRep"].fillna("Unknown").astype(str).str.strip()
df["SalesRep"] = df["SalesRep"].replace("", "Unknown")

df["Channel"] = df["Channel"].astype(str).str.strip().str.title()

# UnitPrice blanks -> average UnitPrice for that Product
avg_price_by_product = df.groupby("Product")["UnitPrice"].transform("mean")
df["UnitPrice"] = df["UnitPrice"].fillna(avg_price_by_product).round(2)

# Quantity blanks -> 1 (documented assumption)
df["Quantity"] = df["Quantity"].fillna(1)

wb = Workbook()
ws = wb.active
ws.title = "Sales_Cleaned"

headers = list(df.columns)
ws.append(headers)
for row in df.itertuples(index=False):
    ws.append(list(row))
n_rows = ws.max_row

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
for c in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=c)
    cell.fill = header_fill
    cell.font = Font(name="Arial", bold=True, color="FFFFFF")
for row in ws.iter_rows(min_row=2, max_row=n_rows, max_col=len(headers)):
    for cell in row:
        cell.font = Font(name="Arial", size=10)
for c in range(1, len(headers) + 1):
    ws.column_dimensions[get_column_letter(c)].width = 14
ws.freeze_panes = "A2"

tbl = Table(displayName="tbl_SalesCleaned", ref=f"A1:{get_column_letter(len(headers))}{n_rows}")
tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(tbl)

# --- CleaningLog sheet ---
log = wb.create_sheet("CleaningLog")
log["A1"] = "Cleaning Log — Day 04"
log["A1"].font = Font(name="Arial", bold=True, size=14)

entries = [
    ("Step", "Action", "Detail"),
    ("1", "Duplicate removal",
     f"Removed {rows_removed} duplicate rows based on OrderID "
     f"({raw_dupes} duplicate OrderIDs detected before cleaning)."),
    ("2", "Missing Region",
     f"{int(raw_missing['Region'])} blank Region values found -> filled with 'Unknown'."),
    ("3", "Missing SalesRep",
     f"{int(raw_missing['SalesRep'])} blank SalesRep values found -> filled with 'Unknown'."),
    ("4", "Missing UnitPrice",
     f"{int(raw_missing['UnitPrice'])} blank UnitPrice values found -> filled with the "
     f"average UnitPrice for that Product (AVERAGEIF equivalent)."),
    ("5", "Missing Quantity",
     f"{int(raw_missing['Quantity'])} blank Quantity values found -> filled with 1 "
     f"(documented assumption: treat as a minimum single-unit order)."),
    ("6", "Text standardization",
     "Region and Channel trimmed of whitespace and standardized to Title Case "
     "(was a mix of UPPER, lower, and padded values)."),
]
r = 3
for row_data in entries:
    for c, val in enumerate(row_data, start=1):
        cell = log.cell(row=r, column=c, value=val)
        if r == 3:
            cell.font = Font(name="Arial", bold=True)
        else:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    r += 1

log.column_dimensions["A"].width = 8
log.column_dimensions["B"].width = 22
log.column_dimensions["C"].width = 90
for rr in range(4, r):
    log.row_dimensions[rr].height = 30

wb.save("/home/claude/30-day-data-journey/phase1-excel/day04/day04_solved.xlsx")
print(f"Day 04 saved. Rows removed: {rows_removed}")
