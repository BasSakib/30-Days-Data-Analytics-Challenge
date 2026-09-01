import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# Use the CLEAN dataset here since Day 02 is about logic, not cleaning
df = pd.read_csv("/home/claude/30-day-data-journey/data/sales_data_clean.csv")

wb = Workbook()
ws = wb.active
ws.title = "Sales"

headers = list(df.columns) + ["OrderSize", "HighValue"]
ws.append(headers)
for row in df.itertuples(index=False):
    ws.append(list(row) + [None, None])

n_rows = ws.max_row
qty_col_letter = "I"   # Quantity
rev_col_letter = "K"   # Revenue
ordersize_col = len(df.columns) + 1
highvalue_col = len(df.columns) + 2

# IF formulas for OrderSize and HighValue
for r in range(2, n_rows + 1):
    ws.cell(row=r, column=ordersize_col).value = (
        f'=IF({qty_col_letter}{r}>=15,"Large",IF({qty_col_letter}{r}>=5,"Medium","Small"))'
    )
    ws.cell(row=r, column=highvalue_col).value = f'=IF({rev_col_letter}{r}>500,"Yes","No")'

# Formatting
header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(name="Arial", bold=True, color="FFFFFF")
for c in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=c)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center")
for row in ws.iter_rows(min_row=2, max_row=n_rows, max_col=len(headers)):
    for cell in row:
        cell.font = Font(name="Arial", size=10)

for c in range(1, len(headers) + 1):
    col_letter = get_column_letter(c)
    ws.column_dimensions[col_letter].width = 14
ws.freeze_panes = "A2"

last_col_letter = get_column_letter(len(headers))
tbl = Table(displayName="tbl_Sales", ref=f"A1:{last_col_letter}{n_rows}")
tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(tbl)

# Summary sheet
summary = wb.create_sheet("Summary")
summary["A1"] = "Sales Summary"
summary["A1"].font = Font(name="Arial", bold=True, size=14)

rows = [
    ("Total Revenue", f"=SUM(Sales!{rev_col_letter}2:{rev_col_letter}{n_rows})", '"$"#,##0.00'),
    ("Average Order Value", f"=AVERAGE(Sales!{rev_col_letter}2:{rev_col_letter}{n_rows})", '"$"#,##0.00'),
    ("Total Quantity Sold", f"=SUM(Sales!{qty_col_letter}2:{qty_col_letter}{n_rows})", "#,##0"),
    ("Number of Orders", f"=COUNTA(Sales!A2:A{n_rows})", "#,##0"),
]
r = 3
for label, formula, numfmt in rows:
    summary.cell(row=r, column=1, value=label).font = Font(name="Arial", bold=True)
    c = summary.cell(row=r, column=2, value=formula)
    c.number_format = numfmt
    r += 1

r += 1
summary.cell(row=r, column=1, value="Order Size Breakdown").font = Font(name="Arial", bold=True, size=12)
r += 1
size_col = get_column_letter(ordersize_col)
for label in ["Large", "Medium", "Small"]:
    summary.cell(row=r, column=1, value=label).font = Font(name="Arial")
    summary.cell(row=r, column=2,
                 value=f'=COUNTIF(Sales!{size_col}2:{size_col}{n_rows},"{label}")')
    r += 1

summary.column_dimensions["A"].width = 24
summary.column_dimensions["B"].width = 16

wb.save("/home/claude/30-day-data-journey/phase1-excel/day02/day02_solved.xlsx")
print("Day 02 saved (pre-recalc).")
