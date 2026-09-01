import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

df = pd.read_csv("/home/claude/30-day-data-journey/data/sales_data_clean.csv")

wb = Workbook()
ws = wb.active
ws.title = "Sales"

headers = list(df.columns) + ["SupplierCode", "CategoryCheck"]
ws.append(headers)
for row in df.itertuples(index=False):
    ws.append(list(row) + [None, None])
n_rows = ws.max_row

product_col_letter = "F"   # Product column in Sales
category_col_letter = "E"  # Category column in Sales
supplier_col = len(df.columns) + 1
catcheck_col = len(df.columns) + 2

# --- ProductRef sheet ---
products = sorted(df["Product"].unique())
prod_to_cat = df.drop_duplicates("Product").set_index("Product")["Category"].to_dict()

ref = wb.create_sheet("ProductRef")
ref.append(["Product", "Category", "SupplierCode"])
for i, p in enumerate(products, start=1):
    ref.append([p, prod_to_cat[p], f"SUP-{i:03d}"])
ref_last_row = ref.max_row

for c in range(1, 4):
    cell = ref.cell(row=1, column=c)
    cell.font = Font(name="Arial", bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
for r in range(2, ref_last_row + 1):
    for c in range(1, 4):
        ref.cell(row=r, column=c).font = Font(name="Arial", size=10)
for col in ["A", "B", "C"]:
    ref.column_dimensions[col].width = 20

# --- Lookup formulas on Sales sheet ---
# XLOOKUP is NOT used (unsupported by this environment's recalculation engine);
# INDEX/MATCH is the reliable, professional equivalent and is shown here.
# A note on real XLOOKUP syntax is included in the Notes below for reference.
for r in range(2, n_rows + 1):
    ws.cell(row=r, column=supplier_col).value = (
        f'=IFERROR(INDEX(ProductRef!$C$2:$C${ref_last_row},'
        f'MATCH({product_col_letter}{r},ProductRef!$A$2:$A${ref_last_row},0)),"Unmapped")'
    )
    ws.cell(row=r, column=catcheck_col).value = (
        f'=VLOOKUP({product_col_letter}{r},ProductRef!$A$2:$C${ref_last_row},2,FALSE)'
    )

# Formatting
header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(name="Arial", bold=True, color="FFFFFF")
for c in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=c)
    cell.fill = header_fill
    cell.font = header_font
for row in ws.iter_rows(min_row=2, max_row=n_rows, max_col=len(headers)):
    for cell in row:
        cell.font = Font(name="Arial", size=10)
for c in range(1, len(headers) + 1):
    ws.column_dimensions[get_column_letter(c)].width = 14
ws.freeze_panes = "A2"

tbl = Table(displayName="tbl_Sales", ref=f"A1:{get_column_letter(len(headers))}{n_rows}")
tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(tbl)

# Notes sheet
notes = wb.create_sheet("Notes")
notes["A1"] = "Lookup Notes — Day 03"
notes["A1"].font = Font(name="Arial", bold=True, size=12)
lines = [
    "",
    "SupplierCode column uses INDEX/MATCH (with IFERROR -> \"Unmapped\" fallback).",
    "In real Excel (Microsoft 365), the equivalent modern formula is:",
    '  =XLOOKUP(F2, ProductRef!A:A, ProductRef!C:C, "Unmapped")',
    "INDEX/MATCH is used in this file for maximum compatibility across Excel versions",
    "and spreadsheet engines, but XLOOKUP is the recommended modern approach if your",
    "organization is fully on Microsoft 365.",
    "",
    "CategoryCheck column uses VLOOKUP to re-confirm Category from ProductRef —",
    "compare this against the original Category column as a sanity check.",
]
for i, line in enumerate(lines, start=2):
    notes.cell(row=i, column=1, value=line).font = Font(name="Arial", size=11)
notes.column_dimensions["A"].width = 90

wb.save("/home/claude/30-day-data-journey/phase1-excel/day03/day03_solved.xlsx")
print("Day 03 saved (pre-recalc).")
