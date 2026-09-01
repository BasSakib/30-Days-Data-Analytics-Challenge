import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import BarChart, LineChart, Reference

df = pd.read_csv("/home/claude/30-day-data-journey/data/sales_data_capstone_final.csv")

wb = Workbook()
ws = wb.active
ws.title = "CapstoneData"

headers = list(df.columns)
ws.append(headers)
for row in df.itertuples(index=False):
    ws.append(list(row))
n_rows = ws.max_row

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
for c in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=c)
    cell.fill = header_fill
    cell.font = Font(name="Arial", bold=True, color="FFFFFF")
for c in range(1, len(headers) + 1):
    ws.column_dimensions[get_column_letter(c)].width = 14
ws.freeze_panes = "A2"
tbl = Table(displayName="tbl_Capstone", ref=f"A1:{get_column_letter(len(headers))}{n_rows}")
tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(tbl)

region_col = get_column_letter(headers.index("Region") + 1)
revenue_col = get_column_letter(headers.index("Revenue") + 1)
category_col = get_column_letter(headers.index("Category") + 1)
product_col = get_column_letter(headers.index("Product") + 1)
month_col = get_column_letter(headers.index("OrderMonth") + 1)
year_col = get_column_letter(headers.index("OrderYear") + 1)

dash = wb.create_sheet("Dashboard", 0)
dash.sheet_view.showGridLines = False
dash["B2"] = "CAPSTONE SALES DASHBOARD"
dash["B2"].font = Font(name="Arial", bold=True, size=18, color="1F4E78")

kpi_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")

total_2024 = df[df["OrderYear"] == 2024]["Revenue"].sum()
total_2025 = df[df["OrderYear"] == 2025]["Revenue"].sum()

kpis = [
    ("Total Revenue", f"=SUM(CapstoneData!{revenue_col}:{revenue_col})", '"$"#,##0'),
    ("Total Orders", f"=COUNTA(CapstoneData!A2:A{n_rows})", "#,##0"),
    ("Average Order Value", f"=SUM(CapstoneData!{revenue_col}:{revenue_col})/COUNTA(CapstoneData!A2:A{n_rows})", '"$"#,##0.00'),
    ("YoY Growth", f"=(SUMIFS(CapstoneData!{revenue_col}:{revenue_col},CapstoneData!{year_col}:{year_col},2025)-SUMIFS(CapstoneData!{revenue_col}:{revenue_col},CapstoneData!{year_col}:{year_col},2024))/SUMIFS(CapstoneData!{revenue_col}:{revenue_col},CapstoneData!{year_col}:{year_col},2024)", "0.0%"),
]
for i, (label, formula, numfmt) in enumerate(kpis):
    col = 2 + i * 3
    label_cell = dash.cell(row=4, column=col, value=label)
    label_cell.font = Font(name="Arial", bold=True, size=10, color="1F4E78")
    val_cell = dash.cell(row=5, column=col, value=formula)
    val_cell.font = Font(name="Arial", bold=True, size=16)
    val_cell.number_format = numfmt
    for rr in (4, 5):
        for cc in (col, col + 1):
            dash.cell(row=rr, column=cc).fill = kpi_fill
    dash.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 1)
    dash.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col + 1)

# Region summary
regions = sorted(df["Region"].unique())
r0 = 8
dash.cell(row=r0, column=1, value="Region").font = Font(bold=True)
dash.cell(row=r0, column=2, value="Revenue").font = Font(bold=True)
for i, reg in enumerate(regions, start=1):
    dash.cell(row=r0 + i, column=1, value=reg)
    dash.cell(row=r0 + i, column=2,
              value=f'=SUMIFS(CapstoneData!{revenue_col}:{revenue_col},CapstoneData!{region_col}:{region_col},A{r0+i})'
              ).number_format = '"$"#,##0'
region_end = r0 + len(regions)

# Category summary
categories = sorted(df["Category"].unique())
c0 = region_end + 3
dash.cell(row=c0, column=1, value="Category").font = Font(bold=True)
dash.cell(row=c0, column=2, value="Revenue").font = Font(bold=True)
for i, cat in enumerate(categories, start=1):
    dash.cell(row=c0 + i, column=1, value=cat)
    dash.cell(row=c0 + i, column=2,
              value=f'=SUMIFS(CapstoneData!{revenue_col}:{revenue_col},CapstoneData!{category_col}:{category_col},A{c0+i})'
              ).number_format = '"$"#,##0'
category_end = c0 + len(categories)

# Monthly trend
months = sorted(df["OrderMonth"].unique())
m0 = category_end + 3
dash.cell(row=m0, column=1, value="Month").font = Font(bold=True)
dash.cell(row=m0, column=2, value="Revenue").font = Font(bold=True)
for i, m in enumerate(months, start=1):
    dash.cell(row=m0 + i, column=1, value=m)
    dash.cell(row=m0 + i, column=2,
              value=f'=SUMIFS(CapstoneData!{revenue_col}:{revenue_col},CapstoneData!{month_col}:{month_col},A{m0+i})'
              ).number_format = '"$"#,##0'
month_end = m0 + len(months)

# Top 10 products table
top10 = df.groupby("Product")["Revenue"].sum().sort_values(ascending=False).head(10)
t0 = month_end + 3
dash.cell(row=t0, column=1, value="Top 10 Products by Revenue").font = Font(bold=True, size=12)
dash.cell(row=t0 + 1, column=1, value="Product").font = Font(bold=True)
dash.cell(row=t0 + 1, column=2, value="Revenue").font = Font(bold=True)
for i, prod in enumerate(top10.index, start=1):
    dash.cell(row=t0 + 1 + i, column=1, value=prod)
    dash.cell(row=t0 + 1 + i, column=2,
              value=f'=SUMIFS(CapstoneData!{revenue_col}:{revenue_col},CapstoneData!{product_col}:{product_col},A{t0+1+i})'
              ).number_format = '"$"#,##0'

# Charts
bar = BarChart()
bar.title = "Revenue by Region"
bar.y_axis.title = "Revenue ($)"
data_ref = Reference(dash, min_col=2, min_row=r0, max_row=region_end)
cats_ref = Reference(dash, min_col=1, min_row=r0 + 1, max_row=region_end)
bar.add_data(data_ref, titles_from_data=True)
bar.set_categories(cats_ref)
bar.height, bar.width = 8, 15
dash.add_chart(bar, "E8")

bar2 = BarChart()
bar2.title = "Revenue by Category"
bar2.y_axis.title = "Revenue ($)"
data_ref2 = Reference(dash, min_col=2, min_row=c0, max_row=category_end)
cats_ref2 = Reference(dash, min_col=1, min_row=c0 + 1, max_row=category_end)
bar2.add_data(data_ref2, titles_from_data=True)
bar2.set_categories(cats_ref2)
bar2.height, bar2.width = 8, 15
dash.add_chart(bar2, "E24")

line = LineChart()
line.title = "Monthly Revenue Trend (24 months)"
line.y_axis.title = "Revenue ($)"
data_ref3 = Reference(dash, min_col=2, min_row=m0, max_row=month_end)
cats_ref3 = Reference(dash, min_col=1, min_row=m0 + 1, max_row=month_end)
line.add_data(data_ref3, titles_from_data=True)
line.set_categories(cats_ref3)
line.height, line.width = 8, 22
dash.add_chart(line, "E40")

note_row = t0 + 13
note_cell = dash.cell(row=note_row, column=1,
    value=("Filter tip: select any cell in tbl_Capstone and use Excel's built-in "
           "table filter dropdowns (or Insert > Slicer via a PivotTable) to filter "
           "by Region or Channel interactively."))
note_cell.font = Font(italic=True, size=9)
note_cell.alignment = Alignment(wrap_text=True)
dash.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=6)

for col, width in [("A", 22), ("B", 14), ("C", 14), ("D", 4), ("E", 12)]:
    dash.column_dimensions[col].width = width

wb.move_sheet("Dashboard", offset=-len(wb.sheetnames))
wb.save("/home/claude/30-day-data-journey/phase5-capstone/day29/capstone_dashboard.xlsx")
print("Capstone dashboard saved (pre-recalc).")
